from openpyxl import load_workbook
wb=load_workbook('PWAMS_Manual_Testing_Report (2).xlsx')
for name in ['Summary','Automation Evidence']:
    ws=wb[name]
    print('='*30,name)
    for r in ws.iter_rows(values_only=True):
        vals=[str(x) for x in r if x is not None]
        if vals: print(' | '.join(vals)[:400])
