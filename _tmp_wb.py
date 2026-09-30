from openpyxl import load_workbook
wb=load_workbook('PWAMS_Manual_Testing_Report (2).xlsx')
print('sheets:',wb.sheetnames)
tot={}
for name in wb.sheetnames:
    ws=wb[name]
    hdr=[c.value for c in ws[1]]
    print('===',name,ws.max_row,ws.max_column,hdr[:14])
