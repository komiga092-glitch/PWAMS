from openpyxl import load_workbook
wb=load_workbook('PWAMS_Manual_Testing_Report (2).xlsx')
want={'AUTH-011','USR-003','DASH-004','DASH-005','PER-008','PER-010','LON-003','LON-009','LON-010','LON-011','AID-008','AID-010','CARE-003','RPT-003','SYNC-002'}
for name in wb.sheetnames:
    if name in ('Summary','Automation Evidence'): continue
    ws=wb[name]
    for r in ws.iter_rows(min_row=2,values_only=True):
        if r[0] in want:
            print('='*20, r[0], name, r[7])
            print('EXPECTED:', r[5])
            print('ACTUAL  :', r[6])
            print('REMARKS :', r[11])
