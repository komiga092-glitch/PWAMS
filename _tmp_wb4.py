from openpyxl import load_workbook
wb=load_workbook('PWAMS_Manual_Testing_Report (2).xlsx')
for name in wb.sheetnames:
    ws=wb[name]
    for r in ws.iter_rows(min_row=2,values_only=True):
        blob=' '.join(str(x) for x in r if x)
        if 'JSON' in blob.upper():
            print(r[0], '|', (r[1] or '')[:50], '|', r[7])
            print('    ACT:', (r[6] or '')[:400])
