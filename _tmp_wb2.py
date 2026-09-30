from openpyxl import load_workbook
wb=load_workbook('PWAMS_Manual_Testing_Report (2).xlsx')
summary=[('Sheet','ID','Status','Scenario')]
for name in wb.sheetnames:
    if name in ('Summary','Automation Evidence'): continue
    ws=wb[name]
    for r in ws.iter_rows(min_row=2,values_only=True):
        if r[0] is None: continue
        summary.append((name,r[0],r[7],(r[1] or '')[:70]))
from collections import Counter
print(Counter(s[2] for s in summary))
for s in summary:
    if s[2] not in ('Pass',):
        print(s)
