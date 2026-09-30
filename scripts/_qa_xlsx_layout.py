"""_qa_xlsx_layout.py - inspect the workbook layout for the report writer."""
from openpyxl import load_workbook

wb = load_workbook("PWAMS_Manual_Testing_Report (2).xlsx")
print("sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print(f"--- {name}: dims={ws.dimensions} max_row={ws.max_row} "
          f"max_col={ws.max_column}")
    for row in ws.iter_rows(min_row=1, max_row=min(6, ws.max_row),
                            max_col=min(12, ws.max_column), values_only=True):
        print("   ", row)
